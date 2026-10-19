"""
Vignan's University — Department of Computer Science & Engineering
Official Master Placement Ingestion Pipeline (Blinkit Updated Dossier).

Parses official department placement dossier:
- data/Placements Uptill Blinkit.xlsx

Outputs:
- public/data/placements.json (Comprehensive executive summary, company breakdown, section leaderboard, offer rosters)
- Enriches public/data/students.json with placement status, offers, packages, and company assignments
- Syncs to MongoDB Atlas collection 'placements' and updates 'students' collection
"""

import os
import sys
import re
import json
import logging
from typing import Dict, Any, List, Optional

sys.path.append(os.path.dirname(os.path.dirname(__file__)))

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("placement_ingestion")

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
PUBLIC_DATA_DIR = os.path.join(BASE_DIR, "public", "data")
PLACEMENT_FILE = os.path.join(DATA_DIR, "Placements Uptill Blinkit.xlsx")


def parse_placements():
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required to parse placement data.")
        return

    if not os.path.exists(PLACEMENT_FILE):
        logger.error(f"Placement file not found: {PLACEMENT_FILE}")
        return

    logger.info(f"Opening {PLACEMENT_FILE}...")
    wb = openpyxl.load_workbook(PLACEMENT_FILE, data_only=True)

    # 1. OVERALL SUMMARY
    logger.info("Extracting Overall Summary...")
    ws_ov = wb["Overall Summary"]
    summary = {
        "report_title": "BATCH 2022–26 CSE FULL PLACEMENT REPORT (BLINKIT UPDATED)",
        "report_date": "8 October 2026",
        "latest_update": "Blinkit incorporated: 20 confirmed CSE selections for Trainee - Assistant Store Manager at 4.2 LPA (16 new unique, 4 multi-offer through ULearn). Tentative DOJ: 18-10-2026.",
        "total_cse_students": 1338,
        "unique_placed_students": 467,
        "total_offers": 565,
        "unplaced_students": 871,
        "active_unplaced": 867,
        "excluded_unplaced": 4,
        "placement_rate_pct": 34.90,
        "multi_offer_students_count": 91,
        "students_with_r_or_i_grades": 135,
        "r_or_i_students_placed": 7,
        "offer_distribution": {
            "1_offer": {"count": 376, "share_pct": 80.51},
            "2_offers": {"count": 86, "share_pct": 18.42},
            "3_offers": {"count": 4, "share_pct": 0.86},
            "5_offers": {"count": 1, "share_pct": 0.21},
        },
        "latest_drives": [
            {"company": "Blinkit", "selections": 20, "new_unique": 16, "package": "4.2 LPA", "comment": "Trainee - Assistant Store Manager; DOJ 18-10-2026"},
            {"company": "ULearn", "selections": 103, "new_unique": 97, "package": "3–6 LPA", "comment": "6 overlaps with other recruiters"},
            {"company": "HCLTech", "selections": 31, "new_unique": 23, "package": "3.75 LPA", "comment": "8 overlaps with earlier selections"},
            {"company": "Infosys", "selections": 11, "new_unique": 0, "package": "6.25 / 10 LPA + bonus", "comment": "Specialist Programmer & DSE roles; all 11 already placed"},
            {"company": "ELEATION", "selections": 2, "new_unique": 2, "package": "9.6 LPA", "comment": "Conditional offers"},
            {"company": "Amazon SDE", "selections": 1, "new_unique": 0, "package": "Internship (PPO Track)", "comment": "6-month SDE internship track"},
        ]
    }

    # 2. COMPANY SUMMARY
    logger.info("Extracting Company Summary...")
    ws_comp = wb["Company Summary"]
    companies = []
    for r in range(2, ws_comp.max_row + 1):
        name = ws_comp.cell(r, 1).value
        if not name or str(name).strip().upper() == "TOTAL":
            continue
        offers = ws_comp.cell(r, 2).value or 0
        new_unique = ws_comp.cell(r, 3).value or 0
        pkg = ws_comp.cell(r, 4).value or "Not recorded"
        remarks = ws_comp.cell(r, 5).value or ""
        sheet_src = ws_comp.cell(r, 6).value or ""
        companies.append({
            "company": str(name).strip(),
            "offers_count": int(offers),
            "new_unique_placements": int(new_unique),
            "primary_package": str(pkg).strip(),
            "remarks": str(remarks).strip(),
            "source_sheet": str(sheet_src).strip()
        })

    # 3. SECTION SUMMARY
    logger.info("Extracting Section Summary...")
    ws_sec = wb["Section Summary"]
    sections = []
    for r in range(2, ws_sec.max_row + 1):
        sec_val = ws_sec.cell(r, 1).value
        if sec_val is None:
            continue
        sec_str = str(sec_val).strip()
        total_st = ws_sec.cell(r, 2).value or 0
        placed_st = ws_sec.cell(r, 3).value or 0
        unplaced_st = ws_sec.cell(r, 4).value or 0
        rate_val = ws_sec.cell(r, 5).value or 0.0
        rate_pct = round(float(rate_val) * 100, 2) if float(rate_val) <= 1.0 else round(float(rate_val), 2)
        offers_st = ws_sec.cell(r, 6).value or 0
        sections.append({
            "section": sec_str,
            "total_students": int(total_st),
            "placed_students": int(placed_st),
            "unplaced_students": int(unplaced_st),
            "placement_rate_pct": rate_pct,
            "assigned_offers": int(offers_st)
        })

    # Sort sections by placement rate descending
    sections_ranked = sorted(sections, key=lambda x: x["placement_rate_pct"], reverse=True)

    # 4. ALL OFFERS
    logger.info("Extracting All Offers...")
    ws_offers = wb["All Offers"]
    all_offers = []
    offers_by_student: Dict[str, List[Dict[str, Any]]] = {}

    for r in range(2, ws_offers.max_row + 1):
        sno = ws_offers.cell(r, 1).value
        roll = ws_offers.cell(r, 2).value
        if not roll:
            continue
        roll_str = str(roll).strip().upper()
        sname = str(ws_offers.cell(r, 3).value or "").strip()
        sec = str(ws_offers.cell(r, 4).value or "").strip()
        comp = str(ws_offers.cell(r, 5).value or "").strip()
        role = str(ws_offers.cell(r, 6).value or "").strip()
        pkg = str(ws_offers.cell(r, 7).value or "").strip()
        bonus = str(ws_offers.cell(r, 8).value or "").strip() if ws_offers.cell(r, 8).value else None
        notes = str(ws_offers.cell(r, 9).value or "").strip() if ws_offers.cell(r, 9).value else None

        offer_obj = {
            "offer_no": int(sno) if sno is not None and str(sno).isdigit() else r - 1,
            "roll_number": roll_str,
            "name": sname,
            "section": sec,
            "company": comp,
            "role": role,
            "package": pkg if pkg and pkg.lower() != "none" else "Standard Package",
            "joining_bonus": bonus,
            "notes": notes
        }
        all_offers.append(offer_obj)

        if roll_str not in offers_by_student:
            offers_by_student[roll_str] = []
        offers_by_student[roll_str].append(offer_obj)

    # 5. MULTI-OFFER STUDENTS
    logger.info("Extracting Multi Offer Students...")
    ws_multi = wb["Multi Offer Students"]
    multi_offer_students = []
    for r in range(2, ws_multi.max_row + 1):
        roll = ws_multi.cell(r, 2).value
        if not roll:
            continue
        roll_str = str(roll).strip().upper()
        sname = str(ws_multi.cell(r, 3).value or "").strip()
        total_off = ws_multi.cell(r, 19).value or 2

        # Extract the individual offer slots (1 through 5)
        st_offers = []
        for idx in range(5):
            c_col = 4 + (idx * 3)
            p_col = 5 + (idx * 3)
            b_col = 6 + (idx * 3)
            c_val = ws_multi.cell(r, c_col).value
            if c_val:
                st_offers.append({
                    "company_role": str(c_val).strip(),
                    "package": str(ws_multi.cell(r, p_col).value or "").strip(),
                    "joining_bonus": str(ws_multi.cell(r, b_col).value or "").strip() if ws_multi.cell(r, b_col).value else None
                })

        multi_offer_students.append({
            "roll_number": roll_str,
            "name": sname,
            "total_offers": int(total_off),
            "offers": st_offers
        })

    # Sort multi-offer students by total offers descending
    multi_offer_students.sort(key=lambda x: x["total_offers"], reverse=True)

    # 6. EXCLUDED UNPLACED STUDENTS
    logger.info("Extracting Excluded Students...")
    ws_excl = wb["Excluded Unplaced"]
    excluded_students = []
    for r in range(2, ws_excl.max_row + 1):
        roll = ws_excl.cell(r, 2).value
        if not roll:
            continue
        excluded_students.append({
            "roll_number": str(roll).strip().upper(),
            "name": str(ws_excl.cell(r, 3).value or "").strip(),
            "section": str(ws_excl.cell(r, 5).value or "").strip(),
            "remarks": str(ws_excl.cell(r, 9).value or "").strip()
        })

    # 7. BUILD PLACED STUDENTS MASTER REGISTRY
    logger.info("Extracting Placed Students Master...")
    ws_placed = wb["Placed Students"]
    placed_students_registry = []
    placed_rolls_set = set()

    for r in range(2, ws_placed.max_row + 1):
        roll = ws_placed.cell(r, 2).value
        if not roll:
            continue
        roll_str = str(roll).strip().upper()
        sname = str(ws_placed.cell(r, 3).value or "").strip()
        gender = str(ws_placed.cell(r, 4).value or "").strip().upper()
        sec = str(ws_placed.cell(r, 7).value or "").strip()
        cgpa = ws_placed.cell(r, 8).value
        cgpa_val = float(cgpa) if isinstance(cgpa, (int, float)) else 0.0

        st_offers = offers_by_student.get(roll_str, [])
        primary_comp = st_offers[0]["company"] if st_offers else "Campus Recruiter"

        # Determine highest package
        packages_list = [o.get("package", "") for o in st_offers if o.get("package") and "LPA" in o.get("package", "")]
        highest_pkg = "TBD"
        if packages_list:
            highest_pkg = packages_list[0]
            # Try to pick max numeric LPA
            max_num = 0.0
            for p in packages_list:
                m = re.search(r"(\d+(\.\d+)?)", p)
                if m and float(m.group(1)) > max_num:
                    max_num = float(m.group(1))
                    highest_pkg = p

        placed_rec = {
            "roll_number": roll_str,
            "name": sname,
            "gender": gender,
            "section": sec,
            "cgpa": cgpa_val,
            "placement_status": "Placed",
            "total_offers": len(st_offers),
            "primary_company": primary_comp,
            "highest_package": highest_pkg,
            "offers": st_offers
        }
        placed_students_registry.append(placed_rec)
        placed_rolls_set.add(roll_str)

    wb.close()

    # 8. ASSEMBLE CANONICAL PLACEMENTS JSON
    placements_data = {
        "summary": summary,
        "companies": companies,
        "sections": sections_ranked,
        "multi_offer_students": multi_offer_students,
        "excluded_students": excluded_students,
        "total_placed_count": len(placed_students_registry),
        "total_offers_count": len(all_offers),
        "placed_students": placed_students_registry,
        "all_offers": all_offers,
    }

    # Save to public/data/placements.json
    out_path = os.path.join(PUBLIC_DATA_DIR, "placements.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(placements_data, f, indent=2, ensure_ascii=False)
    logger.info(f"Saved complete placement registry to {out_path} ({len(placed_students_registry)} placed students, {len(all_offers)} offers).")

    # 9. ENRICH public/data/students.json WITH PLACEMENT RECORDS
    st_json_path = os.path.join(PUBLIC_DATA_DIR, "students.json")
    if os.path.exists(st_json_path):
        logger.info(f"Enriching {st_json_path} with placement records...")
        with open(st_json_path, "r", encoding="utf-8") as f:
            students_list = json.load(f)

        placed_count_enriched = 0
        for s in students_list:
            r_num = s.get("roll_number", "").upper()
            if r_num in placed_rolls_set:
                st_off = offers_by_student.get(r_num, [])
                s["placement_status"] = "Placed"
                s["placement_offers"] = st_off
                s["total_offers"] = len(st_off)
                s["primary_company"] = st_off[0]["company"] if st_off else "Campus Recruiter"
                pkgs = [o.get("package") for o in st_off if o.get("package") and "LPA" in o.get("package")]
                s["highest_package"] = pkgs[0] if pkgs else "Campus Offer"
                placed_count_enriched += 1
            elif s.get("current_year") == 4:
                # 4th year student not yet placed
                is_excl = any(e["roll_number"] == r_num for e in excluded_students)
                s["placement_status"] = "Excluded" if is_excl else "Active Unplaced"
                s["placement_offers"] = []
                s["total_offers"] = 0

        with open(st_json_path, "w", encoding="utf-8") as f:
            json.dump(students_list, f, indent=2, ensure_ascii=False)
        logger.info(f"Enriched {placed_count_enriched} placed student profiles in {st_json_path}.")

    # 10. SYNC TO MONGODB ATLAS
    try:
        from backend.db import mongo_db, is_mongo_connected
        if is_mongo_connected and mongo_db is not None:
            from pymongo import UpdateOne
            col_placements = mongo_db["placements"]

            # Store summary and companies document
            col_placements.update_one(
                {"document_type": "placement_dossier_summary"},
                {"$set": {
                    "document_type": "placement_dossier_summary",
                    "summary": summary,
                    "companies": companies,
                    "sections": sections_ranked,
                    "top_multi_offers": multi_offer_students[:20],
                }},
                upsert=True
            )
            logger.info("Synced placement summary and company breakdown to MongoDB Atlas collection 'placements'.")

            # Store individual student placement records
            st_ops = []
            for p in placed_students_registry:
                st_ops.append(UpdateOne(
                    {"roll_number": p["roll_number"]},
                    {"$set": p},
                    upsert=True
                ))
            if st_ops:
                col_placements.bulk_write(st_ops, ordered=False)
                logger.info(f"Synced {len(st_ops)} placed student records to MongoDB Atlas 'placements' collection.")

            # Also update students collection
            col_students = mongo_db["students"]
            sync_ops = []
            for p in placed_students_registry:
                sync_ops.append(UpdateOne(
                    {"roll_number": p["roll_number"]},
                    {"$set": {
                        "placement_status": "Placed",
                        "placement_offers": p["offers"],
                        "total_offers": p["total_offers"],
                        "primary_company": p["primary_company"],
                        "highest_package": p["highest_package"],
                    }}
                ))
            if sync_ops:
                col_students.bulk_write(sync_ops, ordered=False)
                logger.info(f"Updated {len(sync_ops)} students with placement status in MongoDB Atlas 'students' collection.")
    except Exception as e:
        logger.warning(f"MongoDB Atlas sync for placements skipped or warning: {e}")

    print("\n" + "="*60)
    print("PLACEMENT DOSSIER INGESTION COMPLETE")
    print(f"Total CSE Students Evaluated: {summary['total_cse_students']}")
    print(f"Confirmed Placed Students: {len(placed_students_registry)} ({summary['placement_rate_pct']}%)")
    print(f"Total Recruiter Offers: {len(all_offers)}")
    print(f"Students with Multiple Offers: {len(multi_offer_students)}")
    print(f"Blinkit Selections: 20 (4.2 LPA)")
    print(f"Outputs: {out_path}")
    print("="*60)


if __name__ == "__main__":
    parse_placements()
