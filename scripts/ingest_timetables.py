"""
Vignan's University — Department of Computer Science & Engineering
Official Master Timetable & Faculty Registry Ingestion Pipeline.

Parses official department timetables across Year II, III, and IV:
- data/2nd yt TT (1) (3). 9TH JULY)(3).xlsx
- data/3rd yr UPDATED TT (1) (2) 9TH JULY..xlsx
- data/4th yr TT 17TH JULY(1).xlsx

Parses official faculty registry:
- data/CSE faculty list.xls (157 Official Faculty & Teaching Associates)

Cross-references with:
- public/data/students.json (Mentee assignments per counsellor)

Outputs:
- public/data/timetables.json (Full multi-year section timetables & allocations)
- public/data/faculties.json (Official enriched faculty directory with employee IDs, designations, phones, courses & mentees)
- Syncs to MongoDB Atlas collection 'timetables' and 'faculties'
"""

import os
import sys
import re
import json
import logging
from typing import Dict, Any, List, Set, Optional

sys.path.append(os.path.dirname(os.path.dirname(__file__)))

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("timetable_ingestion")

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
PUBLIC_DATA_DIR = os.path.join(BASE_DIR, "public", "data")

TT_FILES = [
    (2, "2nd yt TT (1) (3). 9TH JULY)(3).xlsx", "2024–2028 Batch"),
    (3, "3rd yr UPDATED TT (1) (2) 9TH JULY..xlsx", "2023–2027 Batch"),
    (4, "4th yr TT 17TH JULY(1).xlsx", "2022–2026 Batch"),
]

FACULTY_XLS_FILE = "CSE faculty list.xls"
DAYS = ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY"]


def normalize_name(name: str) -> str:
    """Normalizes faculty name for robust cross-file fuzzy and token matching."""
    if not name:
        return ""
    n = re.sub(r"^(prof\.|dr\.|mr\.|mrs\.|ms\.)\s*", "", str(name), flags=re.IGNORECASE)
    n = re.sub(r"[\.\,\s\-_]+", "", n).lower()
    return n


def get_name_tokens(name: str) -> Set[str]:
    """Extracts significant name tokens (length >= 4) for cross-matching initials and abbreviations."""
    clean = re.sub(r"^(prof\.|dr\.|mr\.|mrs\.|ms\.)\s*", "", str(name), flags=re.IGNORECASE)
    parts = re.split(r"[\.\,\s\-_]+", clean.lower())
    return {p for p in parts if len(p) >= 4}


def split_faculties(fac_str: str) -> List[Dict[str, str]]:
    """Splits comma-separated faculty string into clean list of name and phone."""
    parts = fac_str.split(",")
    res = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        phone_match = re.search(r"[\(\[](\d{10})[\)\]]", p)
        phone = phone_match.group(1) if phone_match else ""
        clean_name = re.sub(r"[\(\[\{].*?[\)\]\}]", "", p).strip()
        clean_name = re.sub(r"\s+", " ", clean_name)
        clean_name = re.sub(r"^(mon|tue|wed|thu|fri|sat)\s*-\s*", "", clean_name, flags=re.IGNORECASE).strip()
        if len(clean_name) >= 3 and not any(k in clean_name.lower() for k in ["break", "library", "counsel"]):
            res.append({"name": clean_name, "phone": phone})
    return res


def ingest_timetables():
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required to parse timetables")
        return

    try:
        import xlrd
    except ImportError:
        logger.error("xlrd is required to parse faculty list XLS")
        return

    all_timetables = {}
    total_sections_count = 0

    # Intermediate mappings from timetable extraction
    tt_by_phone: Dict[str, Dict[str, Any]] = {}
    tt_by_norm: Dict[str, Dict[str, Any]] = {}
    raw_timetable_faculties: Dict[str, Dict[str, Any]] = {}

    # 1. PARSE SECTION TIMETABLES ACROSS YEARS 2, 3, 4
    for year, fname, batch_desc in TT_FILES:
        fpath = os.path.join(DATA_DIR, fname)
        if not os.path.exists(fpath):
            logger.warning(f"Timetable file not found: {fpath}")
            continue

        logger.info(f"Parsing Year {year} Timetable from {fname}...")
        wb = openpyxl.load_workbook(fpath, data_only=True)
        all_timetables[str(year)] = {}

        for sname in wb.sheetnames:
            sec_num = re.sub(r"[^0-9]", "", sname)
            if not sec_num:
                continue
            sec_key = str(int(sec_num))
            ws = wb[sname]

            section_info = {
                "year": year,
                "batch": batch_desc,
                "section": sec_key,
                "sheet": sname,
                "room": "",
                "class_teacher": "",
                "class_teacher_phone": "",
                "time_slots": [],
                "weekly_schedule": {},
                "course_allocations": [],
            }

            in_allocations = False

            for r in range(1, ws.max_row + 1):
                row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                non_empty = [(c, str(v).strip()) for c, v in enumerate(row_vals, 1) if v is not None and str(v).strip()]
                if not non_empty:
                    continue

                line_str = " ".join(v for c, v in non_empty)

                # Room detection
                if "SECTION" in line_str.upper() and not section_info["room"]:
                    r_match = re.search(r"\[([A-Za-z0-9\-]+)\]|-([0-9]{3})|\(([A-Za-z0-9\-]+)\)", line_str)
                    if r_match:
                        section_info["room"] = r_match.group(1) or r_match.group(2) or r_match.group(3)

                # Day header row
                for c, v in non_empty:
                    if v.upper() == "DAY" and not section_info["time_slots"]:
                        section_info["time_slots"] = [
                            v2 for c2, v2 in non_empty
                            if c2 > c and v2.upper() != "DAY" and not v2.upper().startswith("SEC")
                        ]
                        break

                # Day schedule rows
                for c, v in non_empty:
                    if v.upper() in DAYS:
                        day_name = v.capitalize()
                        periods = [v2 for c2, v2 in non_empty if c2 > c]
                        section_info["weekly_schedule"][day_name] = periods
                        break

                # Class Teacher row
                if "CLASS TEACHER" in line_str.upper():
                    in_allocations = True
                    for c, v in non_empty:
                        if "CLASS TEACHER" not in v.upper():
                            section_info["class_teacher"] = v
                            p_match = re.search(r"[\(\[](\d{10})[\)\]]", v)
                            ct_phone = p_match.group(1) if p_match else ""
                            if ct_phone:
                                section_info["class_teacher_phone"] = ct_phone

                            clean_ct = re.sub(r"[\(\[\{].*?[\)\]\}]", "", v).strip()
                            clean_ct = re.sub(r"\s+", " ", clean_ct)
                            ct_entry = f"Year {year} Sec {sec_key}"

                            if clean_ct:
                                cn = normalize_name(clean_ct)
                                if cn not in tt_by_norm:
                                    tt_by_norm[cn] = {"courses": [], "class_teacher_of": set(), "phones": set()}
                                tt_by_norm[cn]["class_teacher_of"].add(ct_entry)
                                if ct_phone:
                                    tt_by_norm[cn]["phones"].add(ct_phone)
                                    if ct_phone not in tt_by_phone:
                                        tt_by_phone[ct_phone] = {"courses": [], "class_teacher_of": set(), "names": set()}
                                    tt_by_phone[ct_phone]["class_teacher_of"].add(ct_entry)
                                    tt_by_phone[ct_phone]["names"].add(clean_ct)
                            break
                    continue

                # Course allocations
                if in_allocations:
                    if len(non_empty) >= 2:
                        sub = non_empty[0][1]
                        fac = non_empty[1][1]
                        if sub and fac and not any(k in sub.upper() for k in ["NOTE:", "PRINCIPAL", "HOD", "DEAN"]):
                            section_info["course_allocations"].append({
                                "subject": sub,
                                "faculty": fac
                            })

                            fac_list = split_faculties(fac)
                            for f in fac_list:
                                f_name = f["name"]
                                f_phone = f["phone"]
                                c_entry = {"year": year, "section": sec_key, "subject": sub}

                                # Raw faculty fallback
                                if f_name not in raw_timetable_faculties:
                                    raw_timetable_faculties[f_name] = {
                                        "name": f_name,
                                        "phone": f_phone,
                                        "courses": [],
                                        "class_teacher_of": []
                                    }
                                if c_entry not in raw_timetable_faculties[f_name]["courses"]:
                                    raw_timetable_faculties[f_name]["courses"].append(c_entry)
                                if f_phone and not raw_timetable_faculties[f_name].get("phone"):
                                    raw_timetable_faculties[f_name]["phone"] = f_phone

                                # Normalized index
                                fn = normalize_name(f_name)
                                if fn not in tt_by_norm:
                                    tt_by_norm[fn] = {"courses": [], "class_teacher_of": set(), "phones": set()}
                                if c_entry not in tt_by_norm[fn]["courses"]:
                                    tt_by_norm[fn]["courses"].append(c_entry)
                                if f_phone:
                                    tt_by_norm[fn]["phones"].add(f_phone)
                                    if f_phone not in tt_by_phone:
                                        tt_by_phone[f_phone] = {"courses": [], "class_teacher_of": set(), "names": set()}
                                    if c_entry not in tt_by_phone[f_phone]["courses"]:
                                        tt_by_phone[f_phone]["courses"].append(c_entry)
                                    tt_by_phone[f_phone]["names"].add(f_name)

            all_timetables[str(year)][sec_key] = section_info
            total_sections_count += 1

        wb.close()

    # 2. LOAD COUNSELLOR MENTEES FROM STUDENTS.JSON
    st_path = os.path.join(PUBLIC_DATA_DIR, "students.json")
    counsellor_mentees: Dict[str, List[str]] = {}
    if os.path.exists(st_path):
        try:
            with open(st_path, "r", encoding="utf-8") as f:
                students_data = json.load(f)
            for s in students_data:
                c = s.get("counsellor")
                if c:
                    cn = normalize_name(c)
                    if cn not in counsellor_mentees:
                        counsellor_mentees[cn] = []
                    counsellor_mentees[cn].append(s.get("roll_number"))
            logger.info(f"Loaded mentee rosters across {len(counsellor_mentees)} faculty counsellors from students.json.")
        except Exception as e:
            logger.warning(f"Error loading students.json for mentee sync: {e}")

    # 3. PARSE OFFICIAL CSE FACULTY LIST (CSE faculty list.xls)
    fac_xls_path = os.path.join(DATA_DIR, FACULTY_XLS_FILE)
    official_faculties = []
    matched_official_names = set()
    matched_official_phones = set()

    if os.path.exists(fac_xls_path):
        logger.info(f"Ingesting official faculty roster from {fac_xls_path}...")
        xwb = xlrd.open_workbook(fac_xls_path)
        xsheet = xwb.sheet_by_index(0)

        for r in range(3, xsheet.nrows):
            sl_raw = xsheet.cell_value(r, 0)
            emp_raw = xsheet.cell_value(r, 1)
            emp_id = str(int(emp_raw)) if isinstance(emp_raw, float) else str(emp_raw).strip()
            name = str(xsheet.cell_value(r, 2)).strip()
            designation = str(xsheet.cell_value(r, 3)).strip()
            gender = str(xsheet.cell_value(r, 4)).strip().upper()
            mob_raw = xsheet.cell_value(r, 5)
            phone = str(int(mob_raw)) if isinstance(mob_raw, float) else re.sub(r"[^0-9]", "", str(mob_raw))

            is_leadership = any(k in designation.lower() for k in ["dean", "hod"])
            n = normalize_name(name)
            name_tokens = get_name_tokens(name)

            # Match courses and class teacher assignments
            tt_data = None
            if phone and phone in tt_by_phone:
                tt_data = tt_by_phone[phone]
            elif n in tt_by_norm:
                tt_data = tt_by_norm[n]
            else:
                # Token / Substring matching
                for k, v in tt_by_norm.items():
                    if len(k) >= 5 and (k in n or n in k):
                        tt_data = v
                        break
                if not tt_data and name_tokens:
                    for k, v in tt_by_norm.items():
                        k_tokens = set(re.split(r"[\.\,\s\-_]+", k))
                        if len(name_tokens & k_tokens) >= 1 and len(k) >= 6:
                            tt_data = v
                            break

            courses = tt_data["courses"] if tt_data else []
            class_teacher_of = sorted(list(tt_data["class_teacher_of"])) if tt_data else []

            # Match student mentees
            mentees = counsellor_mentees.get(n, [])
            if not mentees:
                for k, v in counsellor_mentees.items():
                    if len(k) >= 5 and (k in n or n in k):
                        mentees = v
                        break
            if not mentees and name_tokens:
                for k, v in counsellor_mentees.items():
                    k_tokens = set(re.split(r"[\.\,\s\-_]+", k))
                    if len(name_tokens & k_tokens) >= 1 and len(k) >= 6:
                        mentees = v
                        break

            email = f"{emp_id}@vignan.ac.in"
            sl_no = int(sl_raw) if isinstance(sl_raw, (int, float)) else r - 2

            fac_record = {
                "sl_no": sl_no,
                "emp_id": emp_id,
                "name": name,
                "designation": designation,
                "gender": gender,
                "phone": phone,
                "email": email,
                "department": "Computer Science & Engineering",
                "is_leadership": is_leadership,
                "is_official_cse": True,
                "class_teacher_of": class_teacher_of,
                "courses": courses,
                "courses_count": len(courses),
                "mentees_count": len(mentees),
                "mentee_rolls": mentees[:15],
            }

            official_faculties.append(fac_record)
            matched_official_names.add(n)
            if phone:
                matched_official_phones.add(phone)

        logger.info(f"Processed {len(official_faculties)} official CSE faculty records.")
    else:
        logger.warning(f"Official faculty list not found at {fac_xls_path}")

    # 4. APPEND ALLIED / INTER-DEPARTMENTAL COURSE INSTRUCTORS
    allied_faculties = []
    for f_name, raw_data in raw_timetable_faculties.items():
        fn = normalize_name(f_name)
        f_phone = raw_data.get("phone", "")
        if fn in matched_official_names or (f_phone and f_phone in matched_official_phones):
            continue
        is_sub = any(len(on) >= 5 and (on in fn or fn in on) for on in matched_official_names)
        if is_sub:
            continue

        allied_faculties.append({
            "sl_no": len(official_faculties) + len(allied_faculties) + 1,
            "emp_id": f"ALLIED-{len(allied_faculties) + 1:03d}",
            "name": f_name,
            "designation": "Course Instructor",
            "gender": "N/A",
            "phone": f_phone,
            "email": "academics@vignan.ac.in",
            "department": "Allied / Inter-Departmental",
            "is_leadership": False,
            "is_official_cse": False,
            "class_teacher_of": raw_data.get("class_teacher_of", []),
            "courses": raw_data.get("courses", []),
            "courses_count": len(raw_data.get("courses", [])),
            "mentees_count": 0,
            "mentee_rolls": [],
        })

    # Sort: Leadership on top, followed by official CSE faculty, then allied instructors
    def fac_sort_key(f):
        if "dean" in f.get("designation", "").lower():
            return (0, 0, f["name"])
        if "hod" in f.get("designation", "").lower() and "deputy" not in f.get("designation", "").lower():
            return (0, 1, f["name"])
        if "deputy" in f.get("designation", "").lower() or "training" in f.get("designation", "").lower():
            return (0, 2, f["name"])
        if f.get("is_official_cse", False):
            return (1, f.get("sl_no", 999), f["name"])
        return (2, 999, f["name"])

    final_faculty_list = sorted(official_faculties + allied_faculties, key=fac_sort_key)

    # Save to public/data/timetables.json
    tt_out_path = os.path.join(PUBLIC_DATA_DIR, "timetables.json")
    with open(tt_out_path, "w", encoding="utf-8") as f:
        json.dump(all_timetables, f, indent=2, ensure_ascii=False)
    logger.info(f"Saved {total_sections_count} section timetables to {tt_out_path}")

    # Save to public/data/faculties.json
    fac_out_path = os.path.join(PUBLIC_DATA_DIR, "faculties.json")
    with open(fac_out_path, "w", encoding="utf-8") as f:
        json.dump(final_faculty_list, f, indent=2, ensure_ascii=False)
    logger.info(f"Saved {len(final_faculty_list)} total faculty records ({len(official_faculties)} official CSE) to {fac_out_path}")

    # 5. MONGODB ATLAS SYNC
    try:
        from backend.db import mongo_db, is_mongo_connected
        if is_mongo_connected and mongo_db is not None:
            from pymongo import UpdateOne
            col_tt = mongo_db["timetables"]
            tt_ops = []
            for y_k, secs in all_timetables.items():
                for s_k, s_data in secs.items():
                    tt_ops.append(UpdateOne({"year": int(y_k), "section": s_k}, {"$set": s_data}, upsert=True))
            if tt_ops:
                col_tt.bulk_write(tt_ops, ordered=False)
                logger.info(f"Synced {len(tt_ops)} timetables to MongoDB Atlas collection 'timetables'.")

            col_fac = mongo_db["faculties"]
            fac_ops = []
            for f in final_faculty_list:
                filter_q = {"emp_id": f["emp_id"]} if f.get("emp_id") else {"name": f["name"]}
                fac_ops.append(UpdateOne(filter_q, {"$set": f}, upsert=True))
            if fac_ops:
                col_fac.bulk_write(fac_ops, ordered=False)
                logger.info(f"Synced {len(fac_ops)} faculties to MongoDB Atlas collection 'faculties'.")
    except Exception as e:
        logger.warning(f"MongoDB sync for timetables/faculties skipped or warning: {e}")

    print("\n" + "="*60)
    print("TIMETABLE & OFFICIAL FACULTY INGESTION COMPLETE")
    print(f"Total Sections Processed: {total_sections_count}")
    print(f"Official CSE Faculty & Teaching Associates: {len(official_faculties)}")
    print(f"Total Unified Faculty Profiles: {len(final_faculty_list)}")
    print(f"Timetables JSON: {tt_out_path}")
    print(f"Faculties JSON: {fac_out_path}")
    print("="*60)


if __name__ == "__main__":
    ingest_timetables()
